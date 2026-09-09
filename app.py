import os
import re
import json
import sqlite3
import hashlib
import zipfile
from datetime import datetime
from functools import wraps

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    session,
    jsonify
)

from werkzeug.utils import secure_filename


# =========================================================
# APP CONFIGURATION
# =========================================================

app = Flask(__name__)

app.secret_key = "facultyflow-ai-secret-key-2026"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DATABASE = os.path.join(
    BASE_DIR,
    "facultyflow.db"
)

UPLOAD_FOLDER = os.path.join(
    BASE_DIR,
    "uploads"
)

os.makedirs(
    UPLOAD_FOLDER,
    exist_ok=True
)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

app.config["MAX_CONTENT_LENGTH"] = (
    25 * 1024 * 1024
)


# =========================================================
# ALLOWED FILE TYPES
# =========================================================

ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "xlsx",
    "pptx",
    "zip",
    "txt",
    "csv"
}


# =========================================================
# DATABASE CONNECTION
# =========================================================

def get_db():

    connection = sqlite3.connect(
        DATABASE
    )

    connection.row_factory = sqlite3.Row

    return connection


# =========================================================
# DATABASE INITIALIZATION
# =========================================================

def init_db():

    connection = get_db()

    # -----------------------------------------------------
    # RECORDS
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT,
            document_type TEXT,
            extracted_data TEXT,
            status TEXT,
            created_at TEXT
        )
    """)

    # -----------------------------------------------------
    # DOCUMENTS
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT,
            file_hash TEXT,
            document_type TEXT,
            file_path TEXT,
            uploaded_at TEXT
        )
    """)

    # -----------------------------------------------------
    # TASKS
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT,
            description TEXT,
            due_date TEXT,
            priority TEXT,
            status TEXT,
            created_at TEXT
        )
    """)

    # -----------------------------------------------------
    # FACULTY CLASSES
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS faculty_classes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            subject TEXT,
            class_name TEXT,
            semester TEXT,
            section TEXT,
            created_at TEXT
        )
    """)

    # -----------------------------------------------------
    # STUDENTS
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT,
            roll_number TEXT,
            department TEXT,
            semester TEXT,
            section TEXT,
            subject TEXT,
            previous_sem TEXT,
            cia TEXT,
            mid_sem TEXT,
            end_sem TEXT,
            project_marks TEXT,
            assignment_marks TEXT,
            attendance TEXT,
            created_at TEXT
        )
    """)

    # -----------------------------------------------------
    # PROJECTS
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS projects (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_name TEXT,
            project_title TEXT,
            guide TEXT,
            status TEXT,
            created_at TEXT
        )
    """)

    # -----------------------------------------------------
    # CLUB ACTIVITIES
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS club_activities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_name TEXT,
            club_name TEXT,
            activity TEXT,
            date TEXT,
            created_at TEXT
        )
    """)

    # -----------------------------------------------------
    # FACULTY ACTIVITIES
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS faculty_activities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            faculty_name TEXT,
            activity TEXT,
            date TEXT,
            duration TEXT,
            status TEXT,
            created_at TEXT
        )
    """)

    # -----------------------------------------------------
    # FACULTY PROFILE
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS faculty_profile (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            faculty_name TEXT,
            email TEXT,
            department TEXT,
            designation TEXT,
            phone TEXT
        )
    """)

    # -----------------------------------------------------
    # FACULTY ATTENDANCE
    # -----------------------------------------------------

    connection.execute("""
        CREATE TABLE IF NOT EXISTS faculty_attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT,
            status TEXT,
            created_at TEXT
        )
    """)

    connection.commit()

    connection.close()


# =========================================================
# LOGIN REQUIRED DECORATOR
# =========================================================

def login_required(function):

    @wraps(function)
    def wrapper(*args, **kwargs):

        if not session.get("logged_in"):

            return redirect(
                url_for("login")
            )

        return function(
            *args,
            **kwargs
        )

    return wrapper


# =========================================================
# LOGIN
# =========================================================

@app.route(
    "/",
    methods=["GET", "POST"]
)
def login():

    if request.method == "POST":

        email = request.form.get(
            "email",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        ).strip()

        if email and password:

            session["logged_in"] = True

            session["faculty_email"] = email

            return redirect(
                url_for("dashboard")
            )

        flash(
            "Please enter your email and password."
        )

    return render_template(
        "login.html"
    )


# =========================================================
# LOGOUT
# =========================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/dashboard")
@login_required
def dashboard():

    connection = get_db()

    # Total records
    total_records = connection.execute("""
        SELECT COUNT(*)
        FROM records
    """).fetchone()[0]

    # Total documents
    total_documents = connection.execute("""
        SELECT COUNT(*)
        FROM documents
    """).fetchone()[0]

    # Total tasks
    total_tasks = connection.execute("""
        SELECT COUNT(*)
        FROM tasks
    """).fetchone()[0]

    # Pending tasks
    pending_tasks = connection.execute("""
        SELECT COUNT(*)
        FROM tasks
        WHERE status != 'Completed'
    """).fetchone()[0]

    # High priority tasks
    high_priority = connection.execute("""
        SELECT COUNT(*)
        FROM tasks
        WHERE priority IN ('High', 'Critical')
        AND status != 'Completed'
    """).fetchone()[0]

    # Completed tasks
    completed_tasks = connection.execute("""
        SELECT COUNT(*)
        FROM tasks
        WHERE status = 'Completed'
    """).fetchone()[0]

    # Students
    total_students = connection.execute("""
        SELECT COUNT(*)
        FROM students
    """).fetchone()[0]

    # Classes
    total_classes = connection.execute("""
        SELECT COUNT(*)
        FROM faculty_classes
    """).fetchone()[0]

    # Projects
    total_projects = connection.execute("""
        SELECT COUNT(*)
        FROM projects
    """).fetchone()[0]

    # Activities
    total_activities = connection.execute("""
        SELECT COUNT(*)
        FROM faculty_activities
    """).fetchone()[0]

    # Attendance
    total_attendance = connection.execute("""
        SELECT COUNT(*)
        FROM faculty_attendance
    """).fetchone()[0]

    # Recent records
    recent_records = connection.execute("""
        SELECT *
        FROM records
        ORDER BY id DESC
        LIMIT 5
    """).fetchall()

    # Recent tasks
    recent_tasks = connection.execute("""
        SELECT *
        FROM tasks
        ORDER BY id DESC
        LIMIT 5
    """).fetchall()

    connection.close()

    return render_template(
        "dashboard.html",

        total_records=total_records,
        total_documents=total_documents,

        total_tasks=total_tasks,
        pending_tasks=pending_tasks,
        high_priority=high_priority,
        completed_tasks=completed_tasks,

        total_students=total_students,
        total_classes=total_classes,
        total_projects=total_projects,
        total_activities=total_activities,
        total_attendance=total_attendance,

        recent_records=recent_records,
        recent_tasks=recent_tasks
    )


# =========================================================
# FILE VALIDATION
# =========================================================

def allowed_file(filename):

    if not filename:
        return False

    if "." not in filename:
        return False

    extension = filename.rsplit(
        ".",
        1
    )[1].lower()

    return extension in ALLOWED_EXTENSIONS


# =========================================================
# FILE HASH
# =========================================================

def calculate_file_hash(filepath):

    sha256 = hashlib.sha256()

    with open(
        filepath,
        "rb"
    ) as file:

        while True:

            chunk = file.read(8192)

            if not chunk:
                break

            sha256.update(chunk)

    return sha256.hexdigest()


# =========================================================
# EXTRACT TEXT FROM FILE
# =========================================================

def extract_text_from_file(filepath):

    if not os.path.exists(filepath):
        return ""

    extension = filepath.rsplit(
        ".",
        1
    )[1].lower()

    # -----------------------------------------------------
    # TXT / CSV
    # -----------------------------------------------------

    if extension in {"txt", "csv"}:

        try:

            with open(
                filepath,
                "r",
                encoding="utf-8",
                errors="ignore"
            ) as file:

                return file.read()

        except Exception as error:

            return (
                "Text extraction error: "
                + str(error)
            )

    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------

    if extension == "pdf":

        try:

            from pypdf import PdfReader

            reader = PdfReader(filepath)

            pages = []

            for page in reader.pages:

                text = page.extract_text()

                if text:
                    pages.append(text)

            return "\n".join(pages)

        except Exception as error:

            return (
                "PDF extraction error: "
                + str(error)
            )

    # -----------------------------------------------------
    # DOCX
    # -----------------------------------------------------

    if extension == "docx":

        try:

            from docx import Document

            document = Document(filepath)

            text = []

            for paragraph in document.paragraphs:

                value = paragraph.text.strip()

                if value:
                    text.append(value)

            for table in document.tables:

                for row in table.rows:

                    values = []

                    for cell in row.cells:

                        cell_value = (
                            cell.text.strip()
                        )

                        if cell_value:

                            values.append(
                                cell_value
                            )

                    if values:

                        text.append(
                            " | ".join(values)
                        )

            return "\n".join(text)

        except Exception as error:

            return (
                "DOCX extraction error: "
                + str(error)
            )

    # -----------------------------------------------------
    # EXCEL
    # -----------------------------------------------------

    if extension == "xlsx":

        try:

            import openpyxl

            workbook = openpyxl.load_workbook(
                filepath,
                data_only=True
            )

            text = []

            for sheet in workbook.worksheets:

                text.append(
                    "Sheet: " + sheet.title
                )

                for row in sheet.iter_rows(
                    values_only=True
                ):

                    values = []

                    for value in row:

                        if value is not None:

                            values.append(
                                str(value)
                            )

                    if values:

                        text.append(
                            " | ".join(values)
                        )

            return "\n".join(text)

        except Exception as error:

            return (
                "Excel extraction error: "
                + str(error)
            )

    # -----------------------------------------------------
    # POWERPOINT
    # -----------------------------------------------------

    if extension == "pptx":

        try:

            from pptx import Presentation

            presentation = Presentation(
                filepath
            )

            text = []

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1
            ):

                text.append(
                    f"Slide {slide_number}"
                )

                for shape in slide.shapes:

                    if hasattr(shape, "text"):

                        value = shape.text.strip()

                        if value:

                            text.append(value)

            return "\n".join(text)

        except Exception as error:

            return (
                "PowerPoint extraction error: "
                + str(error)
            )

    # -----------------------------------------------------
    # ZIP
    # -----------------------------------------------------

    if extension == "zip":

        try:

            text = []

            with zipfile.ZipFile(
                filepath,
                "r"
            ) as archive:

                for filename in archive.namelist():

                    lower_name = filename.lower()

                    if lower_name.endswith(
                        (".txt", ".csv")
                    ):

                        try:

                            content = archive.read(
                                filename
                            ).decode(
                                "utf-8",
                                errors="ignore"
                            )

                            text.append(
                                f"File: {filename}"
                            )

                            text.append(content)

                        except Exception:
                            pass

            return "\n".join(text)

        except Exception as error:

            return (
                "ZIP extraction error: "
                + str(error)
            )

    return ""


# =========================================================
# DOCUMENT TYPE DETECTION
# =========================================================

def detect_document_type(
    text,
    filename=""
):

    content = (
        (text or "")
        + " "
        + (filename or "")
    ).lower()

    # Student
    student_words = [
        "student name",
        "roll number",
        "marks",
        "cia",
        "mid-sem",
        "mid sem",
        "end-sem",
        "end sem",
        "assignment marks",
        "project marks",
        "student academic"
    ]

    if any(
        word in content
        for word in student_words
    ):

        return "Student Academic Record"

    # Faculty
    faculty_words = [
        "faculty name",
        "faculty development",
        "fdp",
        "faculty activity",
        "professional activity",
        "faculty member",
        "workshop",
        "seminar"
    ]

    if any(
        word in content
        for word in faculty_words
    ):

        return "Faculty Activity Document"

    # Attendance
    attendance_words = [
        "attendance",
        "present",
        "absent",
        "leave",
        "attendance percentage"
    ]

    if any(
        word in content
        for word in attendance_words
    ):

        return "Attendance Record"

    # Project
    project_words = [
        "project title",
        "project report",
        "project guide",
        "project supervisor"
    ]

    if any(
        word in content
        for word in project_words
    ):

        return "Project Document"

    # Event
    event_words = [
        "hackathon",
        "event",
        "venue",
        "organized",
        "organised",
        "competition",
        "workshop event",
        "seminar event"
    ]

    if any(
        word in content
        for word in event_words
    ):

        return "Event / Activity Document"

    return "General Document"


# =========================================================
# DYNAMIC FIELD EXTRACTION
# =========================================================

def extract_fields(text):

    fields = {}

    if not text:
        return fields

    # -----------------------------------------------------
    # LABEL : VALUE
    # LABEL = VALUE
    # -----------------------------------------------------

    for line in text.splitlines():

        line = line.strip()

        if not line:
            continue

        match = re.match(
            r"^\s*([^:=]{2,60})\s*[:=]\s*(.+?)\s*$",
            line
        )

        if match:

            label = match.group(1).strip()

            value = match.group(2).strip()

            if label and value:

                label = re.sub(
                    r"\s+",
                    " ",
                    label
                )

                fields[label] = value

    # -----------------------------------------------------
    # EMAIL
    # -----------------------------------------------------

    emails = re.findall(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text
    )

    if emails:

        fields["Email"] = emails[0]

    # -----------------------------------------------------
    # PHONE
    # -----------------------------------------------------

    phones = re.findall(
        r"\b(?:\+91[-\s]?)?[6-9]\d{9}\b",
        text
    )

    if phones:

        fields["Phone"] = phones[0]

    # -----------------------------------------------------
    # PERCENTAGE
    # -----------------------------------------------------

    percentages = re.findall(
        r"\b\d+(?:\.\d+)?\s*%",
        text
    )

    if percentages:

        fields["Percentage"] = percentages[0]

    # -----------------------------------------------------
    # MARKS
    # -----------------------------------------------------

    marks_matches = re.findall(
        r"\b(?:marks|score|grade)\s*[:=]\s*"
        r"([0-9]+(?:\.[0-9]+)?)",
        text,
        flags=re.IGNORECASE
    )

    if marks_matches:

        fields["Marks"] = marks_matches[0]

    # -----------------------------------------------------
    # NUMERIC DATE
    # -----------------------------------------------------

    numeric_dates = re.findall(
        r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b",
        text
    )

    if numeric_dates:

        fields["Date"] = numeric_dates[0]

    else:

        # -------------------------------------------------
        # WRITTEN DATE
        # -------------------------------------------------

        written_dates = re.findall(
            r"\b(?:\d{1,2}\s+)?"
            r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|"
            r"Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|"
            r"Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|"
            r"Nov(?:ember)?|Dec(?:ember)?)"
            r"(?:\s+\d{1,2}(?:st|nd|rd|th)?)?"
            r"(?:,\s*|\s+)?"
            r"(?:20\d{2})?\b",
            text,
            flags=re.IGNORECASE
        )

        if written_dates:

            fields["Date"] = written_dates[0]

    return fields


# =========================================================
# KEYWORD EXTRACTION
# =========================================================

def extract_keywords(text):

    if not text:
        return []

    words = re.findall(
        r"\b[A-Za-z]{4,}\b",
        text.lower()
    )

    stop_words = {
        "this",
        "that",
        "with",
        "from",
        "have",
        "were",
        "will",
        "your",
        "about",
        "which",
        "their",
        "there",
        "faculty",
        "student",
        "document",
        "shall",
        "would",
        "could",
        "should",
        "been",
        "also",
        "into",
        "than",
        "then",
        "where",
        "when",
        "what",
        "these",
        "those",
        "using",
        "used"
    }

    keywords = []

    for word in words:

        if word in stop_words:
            continue

        if word not in keywords:

            keywords.append(word)

    return keywords[:12]


# =========================================================
# SUMMARY
# =========================================================

def generate_summary(text):

    if not text:

        return (
            "No readable text was found "
            "in this document."
        )

    cleaned = " ".join(
        text.split()
    )

    if not cleaned:

        return (
            "No readable text was found "
            "in this document."
        )

    if len(cleaned) <= 400:

        return cleaned

    return cleaned[:400] + "..."


# =========================================================
# DUPLICATE CHECK
# =========================================================

def find_exact_duplicate(file_hash):

    connection = get_db()

    result = connection.execute(
        """
        SELECT *
        FROM documents
        WHERE file_hash = ?
        LIMIT 1
        """,
        (file_hash,)
    ).fetchone()

    connection.close()

    return result


# =========================================================
# SIMILAR FILENAME CHECK
# =========================================================

def find_similar_filename(filename):

    connection = get_db()

    result = connection.execute(
        """
        SELECT *
        FROM documents
        WHERE LOWER(filename) = LOWER(?)
        LIMIT 1
        """,
        (filename,)
    ).fetchone()

    connection.close()

    return result


# =========================================================
# UPLOAD DOCUMENT
# =========================================================

@app.route(
    "/upload",
    methods=["POST"]
)
@login_required
def upload_document():

    # Accept both names
    uploaded_file = request.files.get("file")

    if uploaded_file is None:

        uploaded_file = request.files.get(
            "document"
        )

    # No file
    if uploaded_file is None:

        flash("No file selected.")

        return redirect(
            url_for("documents")
        )

    if uploaded_file.filename == "":

        flash("No file selected.")

        return redirect(
            url_for("documents")
        )

    # Validate
    if not allowed_file(
        uploaded_file.filename
    ):

        flash(
            "Unsupported file type. "
            "Please upload PDF, DOCX, XLSX, "
            "PPTX, ZIP, TXT or CSV."
        )

        return redirect(
            url_for("documents")
        )

    original_filename = (
        uploaded_file.filename
    )

    safe_filename = secure_filename(
        original_filename
    )

    if not safe_filename:

        flash("Invalid filename.")

        return redirect(
            url_for("documents")
        )

    # Unique filename
    timestamp = datetime.now().strftime(
        "%Y%m%d%H%M%S%f"
    )

    stored_filename = (
        timestamp
        + "_"
        + safe_filename
    )

    filepath = os.path.join(
        app.config["UPLOAD_FOLDER"],
        stored_filename
    )

    # Save
    uploaded_file.save(filepath)

    # Hash
    file_hash = calculate_file_hash(
        filepath
    )

    # Exact duplicate
    duplicate = find_exact_duplicate(
        file_hash
    )

    if duplicate:

        try:
            os.remove(filepath)
        except Exception:
            pass

        return render_template(
            "duplicate.html",
            filename=original_filename,
            duplicate=duplicate
        )

    # Similar filename
    similar_filename = find_similar_filename(
        original_filename
    )

    # Extract
    extracted_text = extract_text_from_file(
        filepath
    )

    # Detect type
    document_type = detect_document_type(
        extracted_text,
        original_filename
    )

    # Fields
    extracted_fields = extract_fields(
        extracted_text
    )

    # Keywords
    keywords = extract_keywords(
        extracted_text
    )

    # Summary
    summary = generate_summary(
        extracted_text
    )

    # Warnings
    warnings = []

    if not extracted_text.strip():

        warnings.append(
            "No readable text was extracted."
        )

    if not extracted_fields:

        warnings.append(
            "No labelled fields were detected."
        )

    if similar_filename:

        warnings.append(
            "A document with the same filename "
            "already exists."
        )

    # Database
    connection = get_db()

    try:

        # Save document
        connection.execute(
            """
            INSERT INTO documents
            (
                filename,
                file_hash,
                document_type,
                file_path,
                uploaded_at
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                original_filename,
                file_hash,
                document_type,
                filepath,
                datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            )
        )

        # Save record
        connection.execute(
            """
            INSERT INTO records
            (
                filename,
                document_type,
                extracted_data,
                status,
                created_at
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                original_filename,
                document_type,
                json.dumps(
                    extracted_fields,
                    ensure_ascii=False
                ),
                "Analyzed",
                datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            )
        )

        connection.commit()

    except Exception:

        connection.rollback()

        connection.close()

        raise

    connection.close()

    analysis = {

        "filename": original_filename,

        "document_type": document_type,

        "summary": summary,

        "keywords": keywords,

        "extracted_fields": extracted_fields,

        "extracted_text": extracted_text,

        "warnings": warnings,

        "saved": True
    }

    return render_template(
        "analysis.html",
        analysis=analysis
    )


# =========================================================
# DOCUMENTS PAGE
# =========================================================

@app.route("/documents")
@login_required
def documents():

    connection = get_db()

    documents_data = connection.execute(
        """
        SELECT *
        FROM documents
        ORDER BY id DESC
        """
    ).fetchall()

    connection.close()

    return render_template(
        "documents.html",
        documents=documents_data
    )


# =========================================================
# RECORDS PAGE
# =========================================================

@app.route("/records")
@login_required
def records():

    connection = get_db()

    records_data = connection.execute(
        """
        SELECT *
        FROM records
        ORDER BY id DESC
        """
    ).fetchall()

    connection.close()

    records_list = []

    for record in records_data:

        item = dict(record)

        try:

            item["extracted_data"] = json.loads(
                record["extracted_data"] or "{}"
            )

        except Exception:

            item["extracted_data"] = {}

        records_list.append(item)

    total_records = len(
        records_list
    )

    document_types = len(
        set(
            item.get("document_type")
            for item in records_list
            if item.get("document_type")
        )
    )

    return render_template(
        "records.html",
        records=records_list,
        total_records=total_records,
        document_types=document_types
    )


# =========================================================
# SAVE RECORD
# =========================================================

@app.route(
    "/save-record",
    methods=["POST"]
)
@login_required
def save_record():

    filename = request.form.get(
        "filename",
        "Manual Record"
    ).strip()

    document_type = request.form.get(
        "document_type",
        "General Document"
    ).strip()

    extracted_data = {}

    for key, value in request.form.items():

        if key in {
            "filename",
            "document_type"
        }:
            continue

        if value.strip():

            extracted_data[key] = (
                value.strip()
            )

    connection = get_db()

    connection.execute(
        """
        INSERT INTO records
        (
            filename,
            document_type,
            extracted_data,
            status,
            created_at
        )
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            filename,
            document_type,
            json.dumps(
                extracted_data,
                ensure_ascii=False
            ),
            "Saved",
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )
    )

    connection.commit()

    connection.close()

    flash(
        "Record saved successfully."
    )

    return redirect(
        url_for("records")
    )


# =========================================================
# COMPATIBILITY SAVE RECORD
# =========================================================

@app.route(
    "/save_record",
    methods=["POST"]
)
@login_required
def save_record_old():

    return save_record()


# =========================================================
# TASKS PAGE
# =========================================================

@app.route("/tasks")
@login_required
def tasks():

    connection = get_db()

    rows = connection.execute(
        """
        SELECT *
        FROM tasks
        ORDER BY id DESC
        """
    ).fetchall()

    connection.close()

    task_list = []

    for row in rows:

        task = dict(row)

        # Names expected by tasks.html
        task["name"] = task.get(
            "title",
            ""
        )

        task["deadline"] = task.get(
            "due_date",
            ""
        )

        task["importance"] = task.get(
            "priority",
            "Normal"
        )

        # Display priority
        if task.get("status") == "Completed":

            task["priority"] = "Completed"

            task["priority_class"] = "low"

        elif task.get("priority") == "Critical":

            task["priority"] = "HIGH"

            task["priority_class"] = "high"

        elif task.get("priority") == "Important":

            task["priority"] = "MEDIUM"

            task["priority_class"] = "medium"

        else:

            task["priority"] = "LOW"

            task["priority_class"] = "low"

        task_list.append(task)

    return render_template(
        "tasks.html",
        tasks=task_list
    )


# =========================================================
# ADD TASK
# =========================================================

@app.route(
    "/add_task",
    methods=["POST"]
)
@login_required
def add_task():

    task_name = request.form.get(
        "task",
        ""
    ).strip()

    deadline = request.form.get(
        "deadline",
        ""
    ).strip()

    importance = request.form.get(
        "importance",
        "Normal"
    ).strip()

    if not task_name:

        flash(
            "Please enter a task name."
        )

        return redirect(
            url_for("tasks")
        )

    if not deadline:

        flash(
            "Please select a deadline."
        )

        return redirect(
            url_for("tasks")
        )

    connection = get_db()

    connection.execute(
        """
        INSERT INTO tasks
        (
            title,
            description,
            due_date,
            priority,
            status,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            task_name,
            "",
            deadline,
            importance,
            "Pending",
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )
    )

    connection.commit()

    connection.close()

    flash(
        "Task added successfully."
    )

    return redirect(
        url_for("tasks")
    )


# =========================================================
# COMPLETE TASK
# =========================================================

@app.route(
    "/complete_task/<int:task_id>",
    methods=["POST"]
)
@login_required
def complete_task(task_id):

    connection = get_db()

    connection.execute(
        """
        UPDATE tasks
        SET status = 'Completed'
        WHERE id = ?
        """,
        (task_id,)
    )

    connection.commit()

    connection.close()

    flash(
        "Task marked as completed."
    )

    return redirect(
        url_for("tasks")
    )


# =========================================================
# DELETE TASK
# =========================================================

@app.route(
    "/delete_task/<int:task_id>",
    methods=["POST"]
)
@login_required
def delete_task(task_id):

    connection = get_db()

    connection.execute(
        """
        DELETE FROM tasks
        WHERE id = ?
        """,
        (task_id,)
    )

    connection.commit()

    connection.close()

    flash(
        "Task deleted successfully."
    )

    return redirect(
        url_for("tasks")
    )


# =========================================================
# FACULTY MANAGEMENT PAGE
# =========================================================

@app.route("/faculty-management")
@login_required
def faculty_management():

    connection = get_db()

    profile = connection.execute(
        """
        SELECT *
        FROM faculty_profile
        ORDER BY id DESC
        LIMIT 1
        """
    ).fetchone()

    classes = connection.execute(
        """
        SELECT *
        FROM faculty_classes
        ORDER BY id DESC
        """
    ).fetchall()

    students = connection.execute(
        """
        SELECT *
        FROM students
        ORDER BY id DESC
        """
    ).fetchall()

    projects = connection.execute(
        """
        SELECT *
        FROM projects
        ORDER BY id DESC
        """
    ).fetchall()

    club_activities = connection.execute(
        """
        SELECT *
        FROM club_activities
        ORDER BY id DESC
        """
    ).fetchall()

    faculty_activities = connection.execute(
        """
        SELECT *
        FROM faculty_activities
        ORDER BY id DESC
        """
    ).fetchall()

    attendance = connection.execute(
        """
        SELECT *
        FROM faculty_attendance
        ORDER BY date DESC, id DESC
        """
    ).fetchall()

    connection.close()

    return render_template(
        "faculty_management.html",
        profile=profile,
        classes=classes,
        students=students,
        projects=projects,
        club_activities=club_activities,
        faculty_activities=faculty_activities,
        attendance=attendance
    )


# =========================================================
# SAVE FACULTY PROFILE
# =========================================================

@app.route(
    "/faculty/profile",
    methods=["POST"]
)
@login_required
def save_faculty_profile():

    faculty_name = request.form.get(
        "faculty_name",
        ""
    ).strip()

    email = request.form.get(
        "email",
        ""
    ).strip()

    department = request.form.get(
        "department",
        ""
    ).strip()

    designation = request.form.get(
        "designation",
        ""
    ).strip()

    phone = request.form.get(
        "phone",
        ""
    ).strip()

    connection = get_db()

    existing = connection.execute(
        """
        SELECT id
        FROM faculty_profile
        LIMIT 1
        """
    ).fetchone()

    if existing:

        connection.execute(
            """
            UPDATE faculty_profile
            SET faculty_name = ?,
                email = ?,
                department = ?,
                designation = ?,
                phone = ?
            WHERE id = ?
            """,
            (
                faculty_name,
                email,
                department,
                designation,
                phone,
                existing["id"]
            )
        )

    else:

        connection.execute(
            """
            INSERT INTO faculty_profile
            (
                faculty_name,
                email,
                department,
                designation,
                phone
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                faculty_name,
                email,
                department,
                designation,
                phone
            )
        )

    connection.commit()

    connection.close()

    return redirect(
        url_for("faculty_management")
    )


# =========================================================
# ADD CLASS
# =========================================================

@app.route(
    "/faculty/class/add",
    methods=["POST"]
)
@login_required
def add_class():

    subject = request.form.get(
        "subject",
        ""
    ).strip()

    class_name = request.form.get(
        "class_name",
        ""
    ).strip()

    semester = request.form.get(
        "semester",
        ""
    ).strip()

    section = request.form.get(
        "section",
        ""
    ).strip()

    connection = get_db()

    connection.execute(
        """
        INSERT INTO faculty_classes
        (
            subject,
            class_name,
            semester,
            section,
            created_at
        )
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            subject,
            class_name,
            semester,
            section,
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )
    )

    connection.commit()

    connection.close()

    return redirect(
        url_for("faculty_management")
    )


# =========================================================
# ADD STUDENT
# =========================================================

@app.route(
    "/faculty/student/add",
    methods=["POST"]
)
@login_required
def add_student():

    fields = [
        "name",
        "roll_number",
        "department",
        "semester",
        "section",
        "subject",
        "previous_sem",
        "cia",
        "mid_sem",
        "end_sem",
        "project_marks",
        "assignment_marks",
        "attendance"
    ]

    values = []

    for field in fields:

        values.append(
            request.form.get(
                field,
                ""
            ).strip()
        )

    connection = get_db()

    connection.execute(
        """
        INSERT INTO students
        (
            name,
            roll_number,
            department,
            semester,
            section,
            subject,
            previous_sem,
            cia,
            mid_sem,
            end_sem,
            project_marks,
            assignment_marks,
            attendance,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        values + [
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        ]
    )

    connection.commit()

    connection.close()

    return redirect(
        url_for("faculty_management")
    )


# =========================================================
# ADD PROJECT
# =========================================================

@app.route(
    "/faculty/project/add",
    methods=["POST"]
)
@login_required
def add_project():

    student_name = request.form.get(
        "student_name",
        ""
    ).strip()

    project_title = request.form.get(
        "project_title",
        ""
    ).strip()

    guide = request.form.get(
        "guide",
        ""
    ).strip()

    status = request.form.get(
        "status",
        ""
    ).strip()

    connection = get_db()

    connection.execute(
        """
        INSERT INTO projects
        (
            student_name,
            project_title,
            guide,
            status,
            created_at
        )
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            student_name,
            project_title,
            guide,
            status,
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )
    )

    connection.commit()

    connection.close()

    return redirect(
        url_for("faculty_management")
    )


# =========================================================
# ADD CLUB ACTIVITY
# =========================================================

@app.route(
    "/faculty/club/add",
    methods=["POST"]
)
@login_required
def add_club_activity():

    student_name = request.form.get(
        "student_name",
        ""
    ).strip()

    club_name = request.form.get(
        "club_name",
        ""
    ).strip()

    activity = request.form.get(
        "activity",
        ""
    ).strip()

    date = request.form.get(
        "date",
        ""
    ).strip()

    connection = get_db()

    connection.execute(
        """
        INSERT INTO club_activities
        (
            student_name,
            club_name,
            activity,
            date,
            created_at
        )
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            student_name,
            club_name,
            activity,
            date,
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )
    )

    connection.commit()

    connection.close()

    return redirect(
        url_for("faculty_management")
    )


# =========================================================
# ADD FACULTY ACTIVITY
# =========================================================

@app.route(
    "/faculty/activity/add",
    methods=["POST"]
)
@login_required
def add_faculty_activity():

    faculty_name = request.form.get(
        "faculty_name",
        ""
    ).strip()

    activity = request.form.get(
        "activity",
        ""
    ).strip()

    date = request.form.get(
        "date",
        ""
    ).strip()

    duration = request.form.get(
        "duration",
        ""
    ).strip()

    status = request.form.get(
        "status",
        ""
    ).strip()

    connection = get_db()

    connection.execute(
        """
        INSERT INTO faculty_activities
        (
            faculty_name,
            activity,
            date,
            duration,
            status,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            faculty_name,
            activity,
            date,
            duration,
            status,
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )
    )

    connection.commit()

    connection.close()

    return redirect(
        url_for("faculty_management")
    )


# =========================================================
# ADD FACULTY ATTENDANCE
# =========================================================

@app.route(
    "/faculty/attendance/add",
    methods=["POST"]
)
@login_required
def add_faculty_attendance():

    date = request.form.get(
        "date",
        ""
    ).strip()

    status = request.form.get(
        "status",
        ""
    ).strip()

    connection = get_db()

    connection.execute(
        """
        INSERT INTO faculty_attendance
        (
            date,
            status,
            created_at
        )
        VALUES (?, ?, ?)
        """,
        (
            date,
            status,
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )
    )

    connection.commit()

    connection.close()

    return redirect(
        url_for("faculty_management")
    )


# =========================================================
# AI ASSISTANT PAGE
# =========================================================

@app.route("/ai-assistant")
@login_required
def ai_assistant():

    return render_template(
        "ai_assistant.html"
    )


# =========================================================
# AI ASSISTANT API
# =========================================================

@app.route(
    "/api/analyze-text",
    methods=["POST"]
)
@login_required
def analyze_text():

    data = request.get_json(
        silent=True
    ) or {}

    text = data.get(
        "text",
        ""
    )

    if not isinstance(text, str):

        return jsonify({
            "error": "Invalid text."
        }), 400

    if not text.strip():

        return jsonify({
            "error": "No text provided."
        }), 400

    document_type = detect_document_type(
        text,
        "text.txt"
    )

    fields = extract_fields(
        text
    )

    keywords = extract_keywords(
        text
    )

    summary = generate_summary(
        text
    )

    return jsonify({

        "document_type":
            document_type,

        "summary":
            summary,

        "keywords":
            keywords,

        "extracted_fields":
            fields
    })


# =========================================================
# HEALTH CHECK
# =========================================================

@app.route("/health")
def health():

    return jsonify({

        "status":
            "FacultyFlow AI is running",

        "database":
            os.path.exists(DATABASE)
    })


# =========================================================
# FILE TOO LARGE
# =========================================================

@app.errorhandler(413)
def file_too_large(error):

    flash(
        "File is too large. Maximum size is 25 MB."
    )

    return redirect(
        url_for("documents")
    )


# =========================================================
# PAGE NOT FOUND
# =========================================================

@app.errorhandler(404)
def page_not_found(error):

    return """
    <!DOCTYPE html>
    <html>
    <head>
        <title>FacultyFlow AI</title>
        <style>
            body {
                font-family: Arial, sans-serif;
                background: #f8fafc;
                display: flex;
                justify-content: center;
                align-items: center;
                min-height: 100vh;
            }

            .box {
                background: white;
                padding: 40px;
                border-radius: 15px;
                box-shadow: 0 10px 30px rgba(0,0,0,.1);
                text-align: center;
            }

            a {
                display: inline-block;
                margin-top: 20px;
                padding: 12px 20px;
                background: #4f46e5;
                color: white;
                text-decoration: none;
                border-radius: 8px;
            }
        </style>
    </head>

    <body>

        <div class="box">

            <h1>FacultyFlow AI</h1>

            <p>
                The requested page was not found.
            </p>

            <a href="/dashboard">
                Back to Dashboard
            </a>

        </div>

    </body>
    </html>
    """, 404


# =========================================================
# INTERNAL SERVER ERROR
# =========================================================

@app.errorhandler(500)
def server_error(error):

    return """
    <!DOCTYPE html>
    <html>

    <head>

        <title>FacultyFlow AI - Error</title>

        <style>

            body {
                font-family: Arial, sans-serif;
                background: #f8fafc;
                display: flex;
                justify-content: center;
                align-items: center;
                min-height: 100vh;
            }

            .box {
                background: white;
                padding: 40px;
                border-radius: 15px;
                box-shadow: 0 10px 30px rgba(0,0,0,.1);
                text-align: center;
            }

            a {
                display: inline-block;
                margin-top: 20px;
                padding: 12px 20px;
                background: #4f46e5;
                color: white;
                text-decoration: none;
                border-radius: 8px;
            }

        </style>

    </head>

    <body>

        <div class="box">

            <h1>FacultyFlow AI</h1>

            <p>
                Something went wrong while processing
                the request.
            </p>

            <a href="/">
                Back to Login
            </a>

        </div>

    </body>

    </html>
    """, 500


# =========================================================
# INITIALIZE DATABASE
# =========================================================

init_db()


# =========================================================
# RUN APPLICATION
# =========================================================

if __name__ == "__main__":

    port = int(
        os.environ.get(
            "PORT",
            5000
        )
    )

    host = os.environ.get(
        "HOST",
        "127.0.0.1"
    )

    debug_mode = (
        os.environ.get(
            "FLASK_DEBUG",
            "0"
        ) == "1"
    )

    app.run(
        host=host,
        port=port,
        debug=debug_mode
    )